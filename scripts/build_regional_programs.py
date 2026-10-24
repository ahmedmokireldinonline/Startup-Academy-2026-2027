#!/usr/bin/env python3
import csv, json, re
from pathlib import Path
from collections import Counter
import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter

ROOT=Path(__file__).resolve().parents[1]
RAW=ROOT/'research'/'regional-startup-credits-raw-2026-10-24.json'
OUT_XLSX=ROOT/'Regional_Startup_Credits_Programs_2026-10-24.xlsx'
OUT_CSV=ROOT/'data'/'regional-startup-credits-2026-10-24.csv'
OUT_JSON=ROOT/'data'/'regional-startup-credits-2026-10-24.json'
OUT_CODES=ROOT/'data'/'regional-promo-referral-codes-2026-10-24.csv'
TODAY='2026-10-24'

def rows_from_raw():
    d=json.loads(RAW.read_text(encoding='utf-8'))
    out=[]
    for region in d.get('regions',[]):
        region_name=region.get('region','')
        for i,r in enumerate(region.get('records',[]),1):
            row={
              'Regional Record ID': f"REG-{len(out)+1:03d}",
              'Research Region': region_name,
              'Provider': r.get('provider',''),
              'Program Name': r.get('program_name',''),
              'Country or Region': r.get('country_or_region',''),
              'Benefit Type': r.get('benefit_type',''),
              'Benefit Summary': r.get('benefit_summary',''),
              'Maximum Value': r.get('maximum_value',''),
              'Currency': r.get('currency',''),
              'Program Status': r.get('program_status',''),
              'Eligibility': r.get('eligibility',''),
              'Official Program URL': r.get('official_program_url',''),
              'Application URL': r.get('application_url',''),
              'Promo / Referral Code': r.get('promo_or_referral_code',''),
              'Code Status': r.get('code_status',''),
              'Code Source URL': r.get('code_source_url',''),
              'Last Verified': r.get('last_verified',''),
              'Verification Confidence': r.get('verification_confidence',''),
              'Research Notes': r.get('notes',''),
            }
            row['Has Public Code']='Yes' if row['Promo / Referral Code'] else 'No'
            row['Actionability']='Open/Active' if any(x in row['Program Status'].lower() for x in ['active','open','available','join now']) and not any(x in row['Program Status'].lower() for x in ['closed','finalized','expired','completed','suspended']) else 'Review status'
            out.append(row)
    return out

def add_sheet(wb,title,rows,color='1F4E78'):
    ws=wb.create_sheet(title)
    if not rows:return ws
    headers=list(rows[0])
    ws.append(headers)
    for r in rows: ws.append([r.get(h,'') for h in headers])
    ws.freeze_panes='A2'; ws.auto_filter.ref=ws.dimensions; ws.sheet_properties.tabColor=color
    for c in ws[1]: c.font=Font(bold=True,color='FFFFFF'); c.fill=PatternFill('solid',fgColor=color); c.alignment=Alignment(wrap_text=True,vertical='top')
    for row in ws.iter_rows(min_row=2):
        for c in row:c.alignment=Alignment(wrap_text=True,vertical='top')
    for i,h in enumerate(headers,1):
        n=min(50,max(12,len(str(h)),max((len(str(ws.cell(j,i).value or '')) for j in range(2,min(ws.max_row,25)+1)),default=0)))
        ws.column_dimensions[get_column_letter(i)].width=n
    if ws.max_row>=2:
        ref=f'A1:{get_column_letter(ws.max_column)}{ws.max_row}'
        safe=re.sub(r'[^A-Za-z0-9]','',title)[:20]+'Table'
        t=Table(displayName=safe,ref=ref); t.tableStyleInfo=TableStyleInfo(name='TableStyleMedium2',showRowStripes=True,showColumnStripes=False); ws.add_table(t)
    return ws

def main():
    rows=rows_from_raw()
    codes=[]
    for r in rows:
        if r['Promo / Referral Code'] or r['Code Status']:
            codes.append({'Regional Record ID':r['Regional Record ID'],'Research Region':r['Research Region'],'Provider':r['Provider'],'Program Name':r['Program Name'],'Code Status':r['Code Status'],'Promo / Referral Code':r['Promo / Referral Code'],'Code Source URL':r['Code Source URL'],'Application URL':r['Application URL'],'Last Verified':r['Last Verified'],'Notes':r['Research Notes']})
    summary=[]
    for region in sorted(set(r['Research Region'] for r in rows)):
        rr=[r for r in rows if r['Research Region']==region]
        summary.append({'Research Region':region,'Records':len(rr),'Active/Open-ish':sum(r['Actionability']=='Open/Active' for r in rr),'Review/Closed/Unknown':sum(r['Actionability']!='Open/Active' for r in rr),'Public Codes Found':sum(r['Has Public Code']=='Yes' for r in rr),'High Confidence':sum(r['Verification Confidence']=='High' for r in rr)})
    OUT_CSV.parent.mkdir(parents=True,exist_ok=True); OUT_CODES.parent.mkdir(parents=True,exist_ok=True); (ROOT/'research').mkdir(exist_ok=True)
    for path,data in [(OUT_CSV,rows),(OUT_CODES,codes)]:
        with path.open('w',encoding='utf-8-sig',newline='') as f:
            w=csv.DictWriter(f,fieldnames=list(data[0]) if data else []); w.writeheader(); w.writerows(data)
    OUT_JSON.write_text(json.dumps({'generated_on':TODAY,'counts':{'regions':len(summary),'records':len(rows),'public_codes':sum(r['Has Public Code']=='Yes' for r in rows)},'summary':summary,'programs':rows,'code_registry':codes},ensure_ascii=False,indent=2),encoding='utf-8')
    wb=Workbook(); wb.remove(wb.active)
    add_sheet(wb,'Regional Programs',rows,'1F4E78'); add_sheet(wb,'Promo & Referral Codes',codes,'C65911'); add_sheet(wb,'Regional Summary',summary,'70AD47')
    guide=[
      {'Field':'Research scope','Meaning':'India, Ukraine, Israel, Brazil, Portugal, Spain, EU member states, and Asia, with country-specific rows separated from broader regional rows.'},
      {'Field':'Verification Confidence','Meaning':'High means the researching agent opened and read an official page; Medium means the page or exact terms were only partly verifiable.'},
      {'Field':'Code Status','Meaning':'No public code verified means no provider-published code was found on the verified source; it does not rule out private partner/referral codes.'},
      {'Field':'Maximum Value','Meaning':'Advertised ceiling or stated maximum, not a guaranteed award and not additive across rows.'},
      {'Field':'Actionability','Meaning':'Open/Active is a heuristic based on the reported status; always confirm the live application page before applying.'},
      {'Field':'Research date','Meaning':TODAY},
    ]
    add_sheet(wb,'Field Guide',guide,'A5A5A5'); wb.save(OUT_XLSX)
    print(json.dumps({'records':len(rows),'regions':len(summary),'codes':sum(r['Has Public Code']=='Yes' for r in rows),'xlsx':str(OUT_XLSX)},ensure_ascii=False))
if __name__=='__main__':main()
